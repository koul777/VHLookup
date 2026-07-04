from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples" / "public_admin"
DEMO_OUTPUT = ROOT / "demo_output"
IMAGE_OUTPUT = ROOT / "docs" / "images"


@dataclass(frozen=True)
class TableSpec:
    path: Path
    title: str
    caption: str
    sheet_name: str | None = None


@dataclass(frozen=True)
class FeatureShot:
    title: str
    summary: str
    before: TableSpec
    after: TableSpec
    output_name: str


FEATURE_SHOTS = [
    FeatureShot(
        title="1. 개인정보 마스킹",
        summary="원본 개인정보가 결과 엑셀에서는 값 노출 없이 가려집니다.",
        before=TableSpec(
            SAMPLES / "01_privacy_masking" / "citizen_service_requests.csv",
            "실행 전",
            "원본 CSV",
        ),
        after=TableSpec(
            DEMO_OUTPUT / "01_개인정보_마스킹결과.xlsx",
            "실행 후",
            "마스킹결과",
            "마스킹결과",
        ),
        output_name="feature_01_privacy_masking.png",
    ),
    FeatureShot(
        title="2. 분류별 시트 나누기",
        summary="전체 원본 한 장이 부서별 시트로 나뉩니다.",
        before=TableSpec(
            SAMPLES / "02_split_sheets" / "budget_execution.csv",
            "실행 전",
            "전체 원본 CSV",
        ),
        after=TableSpec(
            DEMO_OUTPUT / "02_분류별_시트나누기.xlsx",
            "실행 후",
            "총무과 시트",
            "총무과",
        ),
        output_name="feature_02_split_sheets.png",
    ),
    FeatureShot(
        title="3. 엑셀/CSV 파일 여러 개 합치기",
        summary="여러 제출 파일을 표준 컬럼 기준으로 하나의 결과표로 합칩니다.",
        before=TableSpec(
            SAMPLES / "03_merge_files" / "row_merge_school_submissions" / "gangbuk_school.csv",
            "실행 전",
            "제출 파일 예시",
        ),
        after=TableSpec(
            DEMO_OUTPUT / "03_제출자료_수합결과.xlsx",
            "실행 후",
            "수합 결과",
            "결과",
        ),
        output_name="feature_03_merge_files.png",
    ),
    FeatureShot(
        title="4. 전/후 파일 검증",
        summary="수정 후 원본만 보면 안 보이는 변경점이 결과 엑셀에서 색과 메모로 표시됩니다.",
        before=TableSpec(
            SAMPLES / "04_before_after_validation" / "payment_after.csv",
            "실행 전",
            "수정 후 원본 CSV",
        ),
        after=TableSpec(
            DEMO_OUTPUT / "04_전후파일_검증결과.xlsx",
            "실행 후",
            "후파일_메모",
            "후파일_메모",
        ),
        output_name="feature_04_before_after_validation.png",
    ),
    FeatureShot(
        title="5. 월별표 목록형 변환",
        summary="옆으로 펼쳐진 월별 열이 열 기준/값 목록형 데이터로 바뀝니다.",
        before=TableSpec(
            SAMPLES / "05_horizontal_table" / "monthly_budget_wide.csv",
            "실행 전",
            "월별 가로표 CSV",
        ),
        after=TableSpec(
            DEMO_OUTPUT / "05_월별표_목록형변환.xlsx",
            "실행 후",
            "목록형 결과",
            "결과",
        ),
        output_name="feature_05_monthly_list.png",
    ),
    FeatureShot(
        title="6. 피벗 요약표 만들기",
        summary="상세 집행 내역이 부서/월 기준 요약표로 집계됩니다.",
        before=TableSpec(
            SAMPLES / "06_pivot_summary" / "budget_execution.csv",
            "실행 전",
            "상세 집행 CSV",
        ),
        after=TableSpec(
            DEMO_OUTPUT / "06_피벗요약표_결과.xlsx",
            "실행 후",
            "피벗요약",
            "피벗요약",
        ),
        output_name="feature_06_pivot_summary.png",
    ),
]


def _font(name: str, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts") / name,
        Path("C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/malgunbd.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


TITLE_FONT = _font("malgunbd.ttf", 30)
SUMMARY_FONT = _font("malgun.ttf", 16)
PANEL_TITLE_FONT = _font("malgunbd.ttf", 20)
CAPTION_FONT = _font("malgun.ttf", 13)
HEADER_FONT = _font("malgunbd.ttf", 15)
BODY_FONT = _font("malgun.ttf", 14)
ARROW_FONT = _font("malgunbd.ttf", 34)


def _rgb_from_fill(fill: PatternFill | None) -> tuple[int, int, int]:
    if fill is None or fill.fill_type != "solid":
        return (255, 255, 255)
    rgb = fill.fgColor.rgb
    if not rgb or rgb == "00000000":
        return (255, 255, 255)
    value = str(rgb)[-6:]
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def _format_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return str(value)


def _read_csv(path: Path, max_rows: int, max_cols: int) -> tuple[list[list[str]], list[list[tuple[int, int, int]]], list[list[bool]]]:
    rows: list[list[str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for index, row in enumerate(reader):
            if index >= max_rows:
                break
            rows.append([_format_value(value) for value in row[:max_cols]])
    if not rows:
        rows = [[""]]
    width = min(max(len(row) for row in rows), max_cols)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    fills = [[(255, 255, 255) for _ in range(width)] for _ in normalized]
    comments = [[False for _ in range(width)] for _ in normalized]
    return normalized, fills, comments


def _read_xlsx(
    path: Path,
    sheet_name: str | None,
    max_rows: int,
    max_cols: int,
) -> tuple[list[list[str]], list[list[tuple[int, int, int]]], list[list[bool]], str]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    row_count = min(sheet.max_row, max_rows)
    col_count = min(sheet.max_column, max_cols)
    rows: list[list[str]] = []
    fills: list[list[tuple[int, int, int]]] = []
    comments: list[list[bool]] = []
    for row in range(1, row_count + 1):
        row_values: list[str] = []
        row_fills: list[tuple[int, int, int]] = []
        row_comments: list[bool] = []
        for col in range(1, col_count + 1):
            cell = sheet.cell(row=row, column=col)
            row_values.append(_format_value(cell.value))
            row_fills.append(_rgb_from_fill(cell.fill))
            row_comments.append(cell.comment is not None)
        rows.append(row_values)
        fills.append(row_fills)
        comments.append(row_comments)
    return rows, fills, comments, sheet.title


def _read_table(spec: TableSpec, max_rows: int, max_cols: int) -> tuple[list[list[str]], list[list[tuple[int, int, int]]], list[list[bool]], str]:
    if spec.path.suffix.lower() == ".csv":
        rows, fills, comments = _read_csv(spec.path, max_rows=max_rows, max_cols=max_cols)
        return rows, fills, comments, spec.path.name
    rows, fills, comments, sheet_title = _read_xlsx(spec.path, spec.sheet_name, max_rows=max_rows, max_cols=max_cols)
    return rows, fills, comments, f"{spec.path.name} / {sheet_title}"


def _clip(draw: ImageDraw.ImageDraw, text: str, width: int, font: ImageFont.ImageFont) -> str:
    if draw.textlength(text, font=font) <= width:
        return text
    clipped = text
    while clipped and draw.textlength(clipped + "...", font=font) > width:
        clipped = clipped[:-1]
    return clipped + "..." if clipped else ""


def _draw_panel(draw: ImageDraw.ImageDraw, spec: TableSpec, x: int, y: int, width: int, height: int) -> None:
    rows, fills, comments, source_label = _read_table(spec, max_rows=11, max_cols=6)
    draw.rounded_rectangle((x, y, x + width, y + height), radius=8, fill=(255, 255, 255), outline=(203, 213, 225), width=1)
    badge_fill = (31, 111, 95) if spec.title == "실행 후" else (71, 85, 105)
    draw.rounded_rectangle((x + 16, y + 14, x + 104, y + 42), radius=6, fill=badge_fill)
    draw.text((x + 28, y + 18), spec.title, fill=(255, 255, 255), font=CAPTION_FONT)
    draw.text((x + 116, y + 15), spec.caption, fill=(23, 50, 77), font=PANEL_TITLE_FONT)
    draw.text((x + 18, y + 48), _clip(draw, source_label, width - 36, CAPTION_FONT), fill=(82, 97, 115), font=CAPTION_FONT)

    table_x = x + 16
    table_y = y + 78
    table_width = width - 32
    row_height = 32
    col_count = max(len(rows[0]), 1)
    col_width = table_width // col_count
    for row_index, row in enumerate(rows):
        cell_y = table_y + row_index * row_height
        for col_index in range(col_count):
            cell_x = table_x + col_index * col_width
            fill = fills[row_index][col_index]
            if row_index == 0:
                fill = (31, 111, 95)
            text_color = (255, 255, 255) if row_index == 0 else (31, 41, 55)
            font = HEADER_FONT if row_index == 0 else BODY_FONT
            draw.rectangle(
                (cell_x, cell_y, cell_x + col_width, cell_y + row_height),
                fill=fill,
                outline=(226, 232, 240),
            )
            text = _clip(draw, row[col_index], col_width - 14, font)
            draw.text((cell_x + 7, cell_y + 8), text, fill=text_color, font=font)
            if comments[row_index][col_index]:
                draw.polygon(
                    [(cell_x + col_width - 12, cell_y), (cell_x + col_width, cell_y), (cell_x + col_width, cell_y + 12)],
                    fill=(220, 38, 38),
                )


def render_feature(shot: FeatureShot) -> Path:
    image_width = 1380
    margin = 24
    gap = 90
    panel_width = (image_width - margin * 2 - gap) // 2
    panel_height = 448
    header_height = 94
    image_height = header_height + panel_height + margin
    image = Image.new("RGB", (image_width, image_height), (244, 247, 250))
    draw = ImageDraw.Draw(image)

    draw.text((margin, 20), shot.title, fill=(23, 50, 77), font=TITLE_FONT)
    draw.text((margin, 60), shot.summary, fill=(82, 97, 115), font=SUMMARY_FONT)

    panel_y = header_height
    before_x = margin
    after_x = margin + panel_width + gap
    _draw_panel(draw, shot.before, before_x, panel_y, panel_width, panel_height)
    arrow_text = "=>"
    arrow_width = draw.textlength(arrow_text, font=ARROW_FONT)
    draw.text(
        (margin + panel_width + (gap - arrow_width) / 2, panel_y + panel_height / 2 - 22),
        arrow_text,
        fill=(31, 111, 95),
        font=ARROW_FONT,
    )
    _draw_panel(draw, shot.after, after_x, panel_y, panel_width, panel_height)

    output_path = IMAGE_OUTPUT / shot.output_name
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path)
    return output_path


def main() -> int:
    missing = sorted({spec.path for shot in FEATURE_SHOTS for spec in (shot.before, shot.after) if not spec.path.exists()})
    if missing:
        joined = ", ".join(str(path.relative_to(ROOT)) for path in missing)
        raise SystemExit(f"required files missing: {joined}. Run scripts/run_public_admin_demo.py first.")
    for shot in FEATURE_SHOTS:
        print(render_feature(shot))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
