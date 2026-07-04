from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from vhlookup_core import (
    AdminWorkbookTools,
    ExcelLoader,
    HeaderDetector,
    PrivacyMaskingEngine,
    ReportWriter,
    SheetDetector,
    WorkbookDiffEngine,
    WorkbookDiffReportWriter,
    template_catalog_frame,
    template_usage_frame,
)
from vhlookup_core.consolidation import ConsolidationEngine
from vhlookup_core.horizontal import HorizontalTableEngine
from vhlookup_core.models import JobResult


SAMPLES = ROOT / "samples" / "public_admin"
OUTPUT = ROOT / "demo_output"


DEMO_TOUR_ROWS = [
    {
        "순서": "01",
        "결과 파일": "01_개인정보_마스킹결과.xlsx",
        "업무 상황": "개인정보 마스킹",
        "열어볼 시트": "마스킹결과, 마스킹내역",
        "확인 포인트": "이름, 고유식별번호, 연락처, 이메일, 계좌번호, 주소, 생년월일 마스킹",
    },
    {
        "순서": "02",
        "결과 파일": "02_분류별_시트나누기.xlsx",
        "업무 상황": "분류별 시트 나누기",
        "열어볼 시트": "결과, 확인사항",
        "확인 포인트": "부서별 시트 자동 생성과 전체 원본 시트 보존",
    },
    {
        "순서": "03",
        "결과 파일": "03_제출자료_수합결과.xlsx",
        "업무 상황": "엑셀/CSV 파일 여러 개 합치기",
        "열어볼 시트": "결과, 확인사항",
        "확인 포인트": "제목 행과 안내문을 건너뛰고 컬럼명을 자동으로 맞춘 통합본",
    },
    {
        "순서": "04",
        "결과 파일": "04_전후파일_검증결과.xlsx",
        "업무 상황": "전/후 파일 검증",
        "열어볼 시트": "후파일_메모, 확인사항",
        "확인 포인트": "값 변경, 추가 행, 누락 행을 색과 메모로 확인",
    },
    {
        "순서": "05",
        "결과 파일": "05_월별표_목록형변환.xlsx",
        "업무 상황": "월별표 목록형 변환",
        "열어볼 시트": "결과",
        "확인 포인트": "월별 컬럼을 열 기준/값 형태의 목록형 자료로 변환",
    },
    {
        "순서": "06",
        "결과 파일": "06_피벗요약표_결과.xlsx",
        "업무 상황": "피벗 요약표 만들기",
        "열어볼 시트": "피벗요약, 확인사항",
        "확인 포인트": "부서별/월별 예산 집행 합계 요약",
    },
]


DEMO_SECURITY_ROWS = [
    {"항목": "AI 사용", "내용": "사용 안 함"},
    {"항목": "외부 API", "내용": "호출 안 함"},
    {"항목": "클라우드 업로드", "내용": "없음"},
    {"항목": "원본 파일", "내용": "수정하지 않음"},
    {"항목": "결과 저장", "내용": "demo_output 폴더에 새 xlsx 파일로 저장"},
    {"항목": "검토 방식", "내용": "첫 시트는 실제 결과, 확인사항 시트는 자동 매칭 근거와 확인 항목"},
    {"항목": "개인정보 점검", "내용": "이름, 연락처, 이메일, 주소, 계좌, 주민등록번호 패턴을 실제 값 저장 없이 유형/건수로 표시"},
]


def load_table(path: Path):
    loader = ExcelLoader()
    header_detector = HeaderDetector()
    sheet_detector = SheetDetector(header_detector)
    source = loader.load(path)
    sheet = sheet_detector.select(source)
    detection = header_detector.detect(sheet)
    return header_detector.apply(sheet, detection)


def attach_auto_summary(result: JobResult, plan, workflow: str) -> JobResult:
    result.mapping_records.extend(plan.evidence_rows)
    result.summary["workflow"] = workflow
    result.summary["auto_key_columns"] = " + ".join(plan.key_spec.reference_key_columns)
    result.summary["auto_target_key_columns"] = " + ".join(plan.key_spec.target_columns())
    if plan.value_columns:
        result.summary["auto_value_columns"] = ", ".join(plan.value_columns)
    return result


def write_demo_tour(output_dir: Path) -> None:
    path = output_dir / "00_샘플_둘러보기.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(DEMO_TOUR_ROWS).to_excel(writer, sheet_name="샘플순서", index=False)
        pd.DataFrame(DEMO_SECURITY_ROWS).to_excel(writer, sheet_name="보안원칙", index=False)
        template_catalog_frame().to_excel(writer, sheet_name="업무템플릿", index=False)
        template_usage_frame().to_excel(writer, sheet_name="명령예시", index=False)
        pd.DataFrame(
            [
                {"단계": 1, "할 일": "00_샘플_둘러보기.xlsx에서 샘플순서 시트를 봅니다."},
                {"단계": 2, "할 일": "03_제출자료_수합결과.xlsx의 확인사항 시트에서 자동 매칭 근거를 봅니다."},
                {"단계": 3, "할 일": "04_전후파일_검증결과.xlsx의 후파일_메모 시트를 봅니다."},
                {"단계": 4, "할 일": "05_월별표_목록형변환.xlsx의 결과 시트를 봅니다."},
                {"단계": 5, "할 일": "06_피벗요약표_결과.xlsx의 피벗요약 시트를 봅니다."},
                {"단계": 6, "할 일": "원본 CSV 파일이 그대로 남아 있는지 확인합니다."},
            ]
        ).to_excel(writer, sheet_name="추천동선", index=False)
        style_workbook(writer.book)


def style_workbook(workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F6F5F")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1 and sheet.max_column > 0:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 48))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[column_letter].width = max(10, min(max_length + 2, 50))


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for old_output in OUTPUT.glob("*.xlsx"):
        old_output.unlink()
    writer = ReportWriter()
    write_demo_tour(OUTPUT)
    tools = AdminWorkbookTools()
    merge_samples = SAMPLES / "03_merge_files"

    PrivacyMaskingEngine().write_xlsx(
        SAMPLES / "01_privacy_masking" / "citizen_service_requests.csv",
        OUTPUT / "01_개인정보_마스킹결과.xlsx",
    )

    tools.write_split_workbook(
        SAMPLES / "02_split_sheets" / "budget_execution.csv",
        OUTPUT / "02_분류별_시트나누기.xlsx",
        split_column="부서",
    )

    consolidated = ConsolidationEngine().consolidate_folder(
        merge_samples / "row_merge_school_submissions",
        template="school_submission_consolidation",
    )
    writer.write_xlsx(consolidated, OUTPUT / "03_제출자료_수합결과.xlsx", include_privacy_scan=False)

    diff = WorkbookDiffEngine().compare_files(
        SAMPLES / "04_before_after_validation" / "payment_before.csv",
        SAMPLES / "04_before_after_validation" / "payment_after.csv",
    )
    WorkbookDiffReportWriter().write_xlsx(diff, OUTPUT / "04_전후파일_검증결과.xlsx")

    horizontal_source = load_table(SAMPLES / "05_horizontal_table" / "monthly_budget_wide.csv")
    horizontal_engine = HorizontalTableEngine()
    detection = horizontal_engine.detect(horizontal_source)
    converted = horizontal_engine.wide_to_long(
        horizontal_source,
        id_columns=[column for column in horizontal_source.columns if column not in detection.value_columns],
        value_columns=detection.value_columns,
    )
    horizontal = JobResult(
        result_frame=converted,
        summary={
            "workflow": "월별표 목록형 변환",
            "row_count": len(converted),
            "auto_value_columns": ", ".join(detection.value_columns),
        },
    )
    writer.write_xlsx(horizontal, OUTPUT / "05_월별표_목록형변환.xlsx")

    tools.write_pivot_workbook(
        SAMPLES / "06_pivot_summary" / "budget_execution.csv",
        OUTPUT / "06_피벗요약표_결과.xlsx",
        row_column="부서",
        column_column="월",
        value_column="금액",
        aggregation="합계",
    )

    for path in sorted(OUTPUT.glob("*.xlsx")):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
