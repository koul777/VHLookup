from __future__ import annotations

import ctypes
import sys
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from tkinter import END, Listbox, Tk, Toplevel, filedialog, messagebox, ttk

from openpyxl import load_workbook

import capture_demo_video_poc as capture


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".tmp" / "demo-video-work"
RAW_DIR = WORK / "raw"
RESULT_DIR = WORK / "results"


@dataclass(frozen=True)
class Scene:
    number: int
    slug: str
    label: str
    intro: str
    input_paths: tuple[Path, ...]
    result_name: str
    result_title: str
    result_description: str
    preferred_sheets: tuple[str, ...]
    result_subtitle: str
    zoom_subtitle: str
    finish_subtitle: str


SAMPLES = ROOT / "samples" / "public_admin"
SCENES = {
    1: Scene(
        number=1,
        slug="privacy_masking",
        label="1. 개인정보 마스킹",
        intro="개인정보 컬럼을 골라 원본 길이 그대로 안전하게 가립니다.",
        input_paths=(SAMPLES / "01_privacy_masking" / "citizen_service_requests.csv",),
        result_name="01_privacy_masking_result.xlsx",
        result_title="마스킹결과 미리보기",
        result_description="자동 인식한 개인정보 컬럼만 별표로 바꾸고 나머지 업무 데이터는 그대로 보존했습니다.",
        preferred_sheets=("마스킹결과", "마스킹확인"),
        result_subtitle="선택한 개인정보는 글자 수를 유지한 별표로 안전하게 가려집니다.",
        zoom_subtitle="원본 데이터 구조와 나머지 업무 컬럼은 그대로 보존됩니다.",
        finish_subtitle="개인정보 마스킹 결과 파일이 완성되었습니다.",
    ),
    2: Scene(
        number=2,
        slug="split_sheets",
        label="2. 분류별 시트 나누기",
        intro="한 장의 취합표를 부서별 시트로 자동 분리합니다.",
        input_paths=(SAMPLES / "02_split_sheets" / "budget_execution.csv",),
        result_name="02_split_sheets_result.xlsx",
        result_title="부서별 시트 결과",
        result_description="부서 열의 값마다 별도 시트를 만들고 전체 자료와 확인사항도 함께 저장했습니다.",
        preferred_sheets=("총무과", "예산과"),
        result_subtitle="총무과·예산과처럼 분류값마다 별도 시트가 자동 생성됩니다.",
        zoom_subtitle="원본 순서를 보존하면서 필요한 부서 자료만 바로 열 수 있습니다.",
        finish_subtitle="분류별 시트 나누기 결과 파일이 완성되었습니다.",
    ),
    3: Scene(
        number=3,
        slug="merge_files",
        label="3. 파일 여러 개 합치기",
        intro="여러 기관의 제출 파일을 컬럼에 맞춰 한 번에 수합합니다.",
        input_paths=(
            SAMPLES / "03_merge_files" / "row_merge_school_submissions" / "gangbuk_school.csv",
            SAMPLES / "03_merge_files" / "row_merge_school_submissions" / "gangnam_school.csv",
        ),
        result_name="03_merge_files_result.xlsx",
        result_title="제출자료 수합결과",
        result_description="두 학교의 같은 양식을 자동으로 판단해 행 방향으로 이어 붙였습니다.",
        preferred_sheets=("결과", "확인사항"),
        result_subtitle="서로 다른 기관 파일의 컬럼을 맞춰 한 결과표로 모읍니다.",
        zoom_subtitle="원본 파일명과 확인사항이 남아 제출 자료를 쉽게 추적할 수 있습니다.",
        finish_subtitle="파일 여러 개 합치기 결과가 완성되었습니다.",
    ),
    4: Scene(
        number=4,
        slug="before_after_validation",
        label="4. 전/후 파일 검증",
        intro="수정 전후의 변경·추가·누락을 색과 메모로 바로 찾습니다.",
        input_paths=(
            SAMPLES / "04_before_after_validation" / "payment_before.csv",
            SAMPLES / "04_before_after_validation" / "payment_after.csv",
        ),
        result_name="04_before_after_validation_result.xlsx",
        result_title="전/후 검증결과",
        result_description="수정 후 표를 기준으로 값 변경, 추가 행, 누락 행을 색과 셀 메모로 표시했습니다.",
        preferred_sheets=("후파일_메모", "확인사항"),
        result_subtitle="노란색 변경 셀과 추가·누락 행을 결과표에서 한눈에 확인합니다.",
        zoom_subtitle="변경 셀의 메모에는 수정 전 값과 수정 후 값이 함께 기록됩니다.",
        finish_subtitle="전/후 파일 검증 결과가 완성되었습니다.",
    ),
    5: Scene(
        number=5,
        slug="horizontal_table",
        label="5. 월별표 목록형 변환",
        intro="옆으로 펼친 월별표를 분석하기 쉬운 세로 목록으로 바꿉니다.",
        input_paths=(SAMPLES / "05_horizontal_table" / "monthly_budget_wide.csv",),
        result_name="05_horizontal_table_result.xlsx",
        result_title="목록형 변환 결과",
        result_description="1월·2월처럼 가로로 펼친 값 열을 ‘열 기준’과 ‘값’ 컬럼의 세로 목록으로 변환했습니다.",
        preferred_sheets=("결과", "확인사항"),
        result_subtitle="월별 열이 반복 가능한 ‘열 기준’과 ‘값’ 행으로 바뀝니다.",
        zoom_subtitle="필터·피벗·차트에 바로 쓸 수 있는 분석용 목록이 됩니다.",
        finish_subtitle="월별표 목록형 변환 결과가 완성되었습니다.",
    ),
    6: Scene(
        number=6,
        slug="pivot_summary",
        label="6. 피벗 요약표 만들기",
        intro="여러 월 시트를 한 번에 합쳐 피벗 요약합니다.",
        input_paths=(SAMPLES / "08_sheet_merge" / "monthly_budget_sheets.xlsx",),
        result_name="06_pivot_summary_result.xlsx",
        result_title="다중 시트 피벗요약",
        result_description="1월과 2월 시트를 함께 읽어 부서별·월별 금액 합계표를 만들었습니다.",
        preferred_sheets=("피벗요약", "확인사항"),
        result_subtitle="선택한 여러 시트의 1월·2월 금액이 하나의 피벗표로 집계됩니다.",
        zoom_subtitle="부서별 월 합계와 전체 합계를 한 표에서 바로 비교할 수 있습니다.",
        finish_subtitle="다중 시트 피벗 요약표가 완성되었습니다.",
    ),
    7: Scene(
        number=7,
        slug="organization_order",
        label="7. 사용자 지정 순서 정렬",
        intro="직제순·우선순위처럼 원하는 값 순서로 전체 행을 정렬합니다.",
        input_paths=(SAMPLES / "07_organization_order" / "department_tasks.csv",),
        result_name="07_organization_order_result.xlsx",
        result_title="사용자 지정 정렬 결과",
        result_description="부서값의 순서를 직접 바꾸고 해당 부서의 전체 행을 같은 순서로 정렬했습니다.",
        preferred_sheets=("결과", "확인사항"),
        result_subtitle="목록에서 정한 부서 순서대로 관련 업무 행이 함께 이동합니다.",
        zoom_subtitle="가나다순이 아닌 직제순·보고순처럼 조직 기준을 그대로 반영합니다.",
        finish_subtitle="사용자 지정 순서 정렬 결과가 완성되었습니다.",
    ),
}


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


def buttons(widget, text: str) -> list[ttk.Button]:
    found: list[ttk.Button] = []
    for child in descendants(widget):
        if isinstance(child, ttk.Button) and str(child.cget("text")) == text:
            found.append(child)
    return found


def first_toplevel(root: Tk, title: str | None = None) -> Toplevel | None:
    candidates = [child for child in root.winfo_children() if isinstance(child, Toplevel)]
    if title is None:
        return candidates[-1] if candidates else None
    for candidate in candidates:
        if title in candidate.title():
            return candidate
    return None


def center_toplevel(root: Tk, top: Toplevel, capture_box: tuple[int, int, int, int]) -> None:
    top.update_idletasks()
    width = min(max(top.winfo_width(), 620), 1240)
    height = min(max(top.winfo_height(), 260), 790)
    source_width = capture_box[2] - capture_box[0]
    source_height = capture_box[3] - capture_box[1]
    x = capture_box[0] + max(0, (source_width - width) // 2)
    y = capture_box[1] + max(0, (source_height - height) // 2 - 10)
    top.geometry(f"{width}x{height}+{x}+{y}")
    top.attributes("-topmost", True)
    top.lift()


def cell_fill(cell) -> str:
    fill = cell.fill
    if fill.fill_type != "solid":
        return "#FFFFFF"
    color = fill.fgColor
    if color.type == "rgb" and color.rgb:
        rgb = str(color.rgb)[-6:]
        if rgb != "000000":
            return f"#{rgb}"
    return "#FFFFFF"


def draw_sheet_table(canvas, worksheet, width: int = 1160, height: int = 370) -> None:
    canvas.delete("all")
    max_columns = min(max(worksheet.max_column, 1), 8)
    max_rows = min(max(worksheet.max_row, 1), 9)
    left, top = 8, 8
    usable_width = width - left * 2
    column_width = usable_width / max_columns
    row_height = min(42, (height - top * 2) / max_rows)

    for row_index in range(1, max_rows + 1):
        for column_index in range(1, max_columns + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            x1 = left + (column_index - 1) * column_width
            y1 = top + (row_index - 1) * row_height
            x2 = x1 + column_width
            y2 = y1 + row_height
            fill = cell_fill(cell)
            if row_index == 1 and fill == "#FFFFFF":
                fill = "#1F6F5F"
            canvas.create_rectangle(x1, y1, x2, y2, fill=fill, outline="#C9D2DC", width=1)
            value = "" if cell.value is None else str(cell.value)
            if len(value) > 20:
                value = value[:18] + "…"
            text_color = "#FFFFFF" if row_index == 1 else "#17324D"
            canvas.create_text(
                (x1 + x2) / 2,
                (y1 + y2) / 2,
                text=value,
                fill=text_color,
                font=("맑은 고딕", 10, "bold" if row_index == 1 else "normal"),
                width=max(40, int(column_width - 12)),
            )


def create_workbook_preview(
    root: Tk,
    scene: Scene,
    result_path: Path,
    capture_box: tuple[int, int, int, int],
    state: capture.CaptureState,
) -> None:
    if not result_path.exists():
        root.after(
            350,
            lambda: create_workbook_preview(root, scene, result_path, capture_box, state),
        )
        return

    workbook = load_workbook(result_path, data_only=False)
    selected_sheets = [name for name in scene.preferred_sheets if name in workbook.sheetnames]
    if not selected_sheets:
        selected_sheets = workbook.sheetnames[:2]

    preview = Toplevel(root)
    preview.title(scene.result_title)
    root_x, root_y = capture_box[0], capture_box[1]
    preview.geometry(f"1280x650+{root_x + 160}+{root_y + 115}")
    preview.transient(root)
    preview.attributes("-topmost", True)
    preview.lift()

    frame = ttk.Frame(preview, padding=20)
    frame.pack(fill="both", expand=True)
    ttk.Label(frame, text=scene.result_title, font=("맑은 고딕", 18, "bold")).pack(anchor="w")
    ttk.Label(
        frame,
        text=scene.result_description,
        foreground="#526173",
        font=("맑은 고딕", 11),
        wraplength=1160,
        justify="left",
    ).pack(anchor="w", pady=(4, 12))

    notebook = ttk.Notebook(frame)
    notebook.pack(fill="both", expand=True)
    first_canvas = None
    for sheet_name in selected_sheets[:2]:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text=sheet_name)
        table_canvas = __import__("tkinter").Canvas(
            tab,
            width=1160,
            height=370,
            background="#FFFFFF",
            highlightthickness=0,
        )
        table_canvas.pack(fill="both", expand=True)
        table_canvas.update_idletasks()
        draw_sheet_table(table_canvas, workbook[sheet_name])
        if first_canvas is None:
            first_canvas = table_canvas

    ttk.Label(
        frame,
        text=f"완료 · 결과 시트 {len(workbook.sheetnames)}개 · {result_path.name}",
        foreground="#1F6F5F",
        font=("맑은 고딕", 11, "bold"),
    ).pack(anchor="w", pady=(12, 0))

    preview.update_idletasks()
    if first_canvas is not None:
        state.update(cursor_target=capture.output_position(first_canvas, capture_box))
    state.update(subtitle=scene.result_subtitle)


def configure_file_choosers(scene: Scene) -> None:
    queue = [str(path) for path in scene.input_paths]

    def choose_one(*_args, **_kwargs) -> str:
        if not queue:
            return str(scene.input_paths[-1])
        return queue.pop(0)

    filedialog.askopenfilename = choose_one
    filedialog.askopenfilenames = lambda *_args, **_kwargs: tuple(str(path) for path in scene.input_paths)
    messagebox.showinfo = lambda *_args, **_kwargs: None
    messagebox.showwarning = lambda *_args, **_kwargs: None
    messagebox.showerror = lambda *_args, **_kwargs: None


def run_scene(scene: Scene) -> int:
    for input_path in scene.input_paths:
        if not input_path.exists():
            raise FileNotFoundError(input_path)

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    result_path = RESULT_DIR / scene.result_name
    raw_video = RAW_DIR / f"{scene.number:02d}_{scene.slug}.webm"
    result_path.unlink(missing_ok=True)
    raw_video.unlink(missing_ok=True)
    capture.RAW_VIDEO = raw_video

    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass

    root = Tk()
    root.geometry("1600x900+60+40")
    root.update_idletasks()
    app = capture.LocalApp(root)
    root.geometry("1600x900+60+40")
    root.update_idletasks()
    root.lift()
    root.attributes("-topmost", True)

    capture_box = (
        root.winfo_rootx(),
        root.winfo_rooty(),
        root.winfo_rootx() + root.winfo_width(),
        root.winfo_rooty() + root.winfo_height(),
    )
    state = capture.CaptureState()
    state.update(label=scene.label, subtitle=scene.intro)
    configure_file_choosers(scene)

    app._ask_file_or_none = lambda _title: scene.input_paths[0]
    app._ask_files_or_none = lambda _title: list(scene.input_paths)
    app._ask_save_path_or_none = lambda _title, _base_name: result_path
    app.open_folder = lambda _path: None

    original_run_with_progress = app._run_with_progress

    def paced_run_with_progress(label, job, *args, **kwargs):
        def paced_job(progress):
            def paced_progress(percent: int, text: str) -> None:
                progress(percent, text)
                time.sleep(0.18)

            return job(paced_progress)

        def center_progress() -> None:
            progress_dialog = first_toplevel(root, "진행률")
            if progress_dialog is not None:
                center_toplevel(root, progress_dialog, capture_box)

        root.after(140, center_progress)
        return original_run_with_progress(label, paced_job, *args, **kwargs)

    app._run_with_progress = paced_run_with_progress

    ready = threading.Event()
    capture_thread = threading.Thread(
        target=capture.capture_loop,
        args=(state, capture_box, ready),
        daemon=True,
    )
    capture_thread.start()
    ready.wait(timeout=10)

    def click(widget, subtitle: str | None = None, delay: int = 950) -> None:
        if subtitle:
            state.update(subtitle=subtitle)
        state.update(cursor_target=capture.output_position(widget, capture_box))

        def invoke() -> None:
            state.update(click_until=time.monotonic() + 0.65)
            widget.invoke()

        root.after(delay, invoke)

    def open_main_menu() -> None:
        main_button = capture.button_with(root, lambda text: text.startswith(f"{scene.number}."))
        if main_button is None:
            raise RuntimeError(f"{scene.number}번 메뉴 버튼을 찾지 못했습니다.")
        click(main_button, f"{scene.number}번 메뉴를 선택해 실제 샘플 파일로 실행합니다.", 1100)
        root.after(1120, automate_dialog)

    def automate_dialog() -> None:
        if scene.number == 1:
            automate_privacy()
        elif scene.number == 2:
            automate_split()
        elif scene.number == 3:
            automate_consolidate()
        elif scene.number == 4:
            automate_diff()
        elif scene.number == 5:
            automate_horizontal()
        elif scene.number == 6:
            automate_pivot()
        elif scene.number == 7:
            automate_order()

    def automate_privacy() -> None:
        dialog = first_toplevel(root, "마스킹할 컬럼 선택")
        if dialog is None:
            root.after(200, automate_privacy)
            return
        center_toplevel(root, dialog, capture_box)
        confirm = buttons(dialog, "이 컬럼들 마스킹")[0]
        click(confirm, "자동 인식된 개인정보 컬럼을 확인하고 실행합니다.", 2200)

    def automate_split() -> None:
        dialog = first_toplevel(root, "분류 기준 열 선택")
        if dialog is None:
            root.after(200, automate_split)
            return
        center_toplevel(root, dialog, capture_box)
        confirm = buttons(dialog, "확인")[0]
        click(confirm, "자동 추천된 ‘부서’ 열을 분류 기준으로 사용합니다.", 2200)

    def automate_consolidate() -> None:
        dialog = first_toplevel(root, "엑셀/CSV 파일 여러 개 합치기")
        if dialog is None:
            root.after(200, automate_consolidate)
            return
        center_toplevel(root, dialog, capture_box)
        upload = buttons(dialog, "파일 올리기")[0]
        click(upload, "두 기관의 제출 파일을 한 번에 올립니다.", 900)

        def run() -> None:
            execute = buttons(dialog, "실행")[0]
            click(execute, "자동 추천된 행 합치기 미리보기를 확인하고 실행합니다.", 900)

        root.after(5200, run)

    def automate_diff() -> None:
        dialog = first_toplevel(root, "전/후 파일 검증")
        if dialog is None:
            root.after(200, automate_diff)
            return
        center_toplevel(root, dialog, capture_box)
        file_buttons = buttons(dialog, "파일 선택")
        click(file_buttons[0], "수정 전 파일과 수정 후 파일을 차례로 선택합니다.", 700)
        root.after(1500, lambda: click(file_buttons[1], None, 500))

        def run() -> None:
            execute = buttons(dialog, "실행")[0]
            click(execute, "미리보기에서 변경 내역을 확인하고 검증을 실행합니다.", 900)

        root.after(4600, run)

    def automate_horizontal() -> None:
        dialog = first_toplevel(root, "월별표 목록형 변환")
        if dialog is None:
            root.after(200, automate_horizontal)
            return
        center_toplevel(root, dialog, capture_box)
        choose = buttons(dialog, "파일 선택")[0]
        click(choose, "월별 값 열을 자동 감지해 목록형 미리보기를 만듭니다.", 800)

        def run() -> None:
            execute = buttons(dialog, "실행")[0]
            click(execute, "자동 감지된 행 기준 컬럼을 확인하고 변환합니다.", 900)

        root.after(4700, run)

    def automate_pivot() -> None:
        dialog = first_toplevel(root, "피벗 요약표 만들기")
        if dialog is None:
            root.after(200, automate_pivot)
            return
        center_toplevel(root, dialog, capture_box)
        choose = buttons(dialog, "파일 선택")[0]
        click(choose, "1월·2월·빈시트가 든 파일에서 처리할 시트를 불러옵니다.", 700)

        def set_options() -> None:
            combo_boxes = [widget for widget in descendants(dialog) if isinstance(widget, ttk.Combobox)]
            for combo in combo_boxes:
                values = [str(value) for value in combo.cget("values")]
                if "합계" in values and "건수" in values:
                    combo.set("합계")
                elif "(행 개수)" in values:
                    combo.set("금액")
                elif "(선택 안 함)" in values:
                    combo.set("월")
                elif "부서" in values:
                    combo.set("부서")
            state.update(subtitle="행은 부서, 열은 월, 값은 금액, 집계는 합계로 지정합니다.")
            refresh = buttons(dialog, "미리보기 새로고침")[0]
            refresh.invoke()

        def run() -> None:
            execute = buttons(dialog, "실행")[0]
            click(execute, "선택한 여러 시트를 함께 집계해 피벗표를 저장합니다.", 900)

        root.after(2700, set_options)
        root.after(5000, run)

    def automate_order() -> None:
        dialog = first_toplevel(root, "사용자 지정 순서 정렬")
        if dialog is None:
            root.after(200, automate_order)
            return
        center_toplevel(root, dialog, capture_box)
        choose = buttons(dialog, "파일 선택")[0]
        click(choose, "정렬할 ‘부서’ 열과 고유값 목록을 불러옵니다.", 700)

        def move_value() -> None:
            order_lists = [widget for widget in descendants(dialog) if isinstance(widget, Listbox)]
            if order_lists and order_lists[0].size() > 1:
                order_list = order_lists[0]
                order_list.selection_clear(0, END)
                order_list.selection_set(order_list.size() - 1)
                order_list.see(order_list.size() - 1)
                move_top = buttons(dialog, "맨 위")[0]
                state.update(subtitle="값을 위·아래로 움직여 원하는 직제순을 직접 만듭니다.")
                move_top.invoke()

        def run() -> None:
            execute = buttons(dialog, "실행")[0]
            click(execute, "지정한 순서대로 관련 업무 행 전체를 정렬합니다.", 900)

        root.after(2500, move_value)
        root.after(5000, run)

    def show_preview() -> None:
        create_workbook_preview(root, scene, result_path, capture_box, state)

    def zoom_result() -> None:
        state.update(
            subtitle=scene.zoom_subtitle,
            zoom_scale=1.35,
            zoom_center=(960.0, 480.0),
        )

    def unzoom_result() -> None:
        state.update(
            subtitle=scene.finish_subtitle,
            zoom_scale=1.0,
            cursor_target=(1480.0, 850.0),
        )

    def stop() -> None:
        state.update(stop=True)
        root.after(300, root.destroy)

    root.after(2200, open_main_menu)
    root.after(13000, show_preview)
    root.after(18300, zoom_result)
    root.after(23200, unzoom_result)
    root.after(30000, stop)

    try:
        root.mainloop()
    finally:
        state.update(stop=True)
        capture_thread.join(timeout=45)

    if not raw_video.exists() or raw_video.stat().st_size < 100_000:
        raise RuntimeError(f"{scene.number}번 메뉴 영상이 생성되지 않았습니다.")
    if not result_path.exists():
        raise RuntimeError(f"{scene.number}번 메뉴 결과 파일이 생성되지 않았습니다.")
    print(raw_video)
    return 0


def main() -> int:
    if len(sys.argv) != 2 or not sys.argv[1].isdigit():
        raise SystemExit("usage: capture_demo_videos.py <scene 1-7>")
    scene_number = int(sys.argv[1])
    if scene_number not in SCENES:
        raise SystemExit("scene must be between 1 and 7")
    return run_scene(SCENES[scene_number])


if __name__ == "__main__":
    raise SystemExit(main())
