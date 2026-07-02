from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vhlookup_core.inspection import InspectionResult


class InspectionReportWriter:
    def write_xlsx(self, result: InspectionResult, path: str | Path) -> Path:
        output_path = Path(path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            self._guide_frame(result).to_excel(writer, sheet_name="먼저확인", index=False)
            result.files.to_excel(writer, sheet_name="파일점검", index=False)
            result.sheets.to_excel(writer, sheet_name="시트헤더점검", index=False)
            if not result.column_comparison.empty:
                result.column_comparison.to_excel(writer, sheet_name="열비교", index=False)
            if not result.column_profiles.empty:
                result.column_profiles.to_excel(writer, sheet_name="컬럼요약", index=False)
            if not result.column_mappings.empty:
                result.column_mappings.to_excel(writer, sheet_name="컬럼매칭점검", index=False)
                self._mapping_evidence_frame(result).to_excel(writer, sheet_name="자동추천근거", index=False)
            if not result.privacy_records.empty:
                result.privacy_records.to_excel(writer, sheet_name="개인정보점검", index=False)
            self._issue_frame(result).to_excel(writer, sheet_name="확인필요", index=False)
            pd.DataFrame(
                [{"항목": key, "값": value} for key, value in result.summary.items()]
            ).to_excel(writer, sheet_name="점검요약", index=False)
            self._style_workbook(writer.book)
        return output_path

    def _guide_frame(self, result: InspectionResult) -> pd.DataFrame:
        rows = [
            {
                "확인 항목": "파일 수",
                "값": result.summary.get("file_count", 0),
                "보는 방법": "파일점검 시트에서 실패 파일이 있는지 확인하세요.",
            },
            {
                "확인 항목": "헤더 자동탐지",
                "값": len(result.sheets),
                "보는 방법": "시트헤더점검 시트에서 헤더 행과 인식 컬럼을 확인하세요.",
            },
            {
                "확인 항목": "파일별 열 차이",
                "값": result.summary.get("unique_column_count", 0),
                "보는 방법": "열비교 시트에서 어떤 파일에 어떤 열이 빠졌는지 확인하세요.",
            },
            {
                "확인 항목": "컬럼별 값 분포",
                "값": len(result.column_profiles),
                "보는 방법": "컬럼요약 시트에서 빈값, 고유값 수, 추정 유형을 확인하세요.",
            },
            {
                "확인 항목": "컬럼 자동매칭",
                "값": len(result.column_mappings),
                "보는 방법": "컬럼매칭점검 시트에서 표준 컬럼과 원본 컬럼 연결을 확인하세요.",
            },
            {
                "확인 항목": "개인정보 의심 컬럼",
                "값": len(result.privacy_records),
                "보는 방법": "개인정보점검 시트가 있으면 공유 전 배포 범위와 필요 여부를 확인하세요.",
            },
            {
                "확인 항목": "확인 필요",
                "값": len(result.issues),
                "보는 방법": "확인필요 시트가 비어 있으면 사전점검 단계에서는 큰 오류가 없습니다.",
            },
        ]
        return pd.DataFrame(rows)

    def _mapping_evidence_frame(self, result: InspectionResult) -> pd.DataFrame:
        frame = result.column_mappings.copy()
        if frame.empty:
            return frame
        frame["확인 안내"] = frame.apply(
            lambda row: (
                "자동 매칭됨"
                if str(row.get("원본 컬럼", "")).strip() and float(row.get("신뢰도", 0) or 0) > 0
                else "원본에서 대응 컬럼을 찾지 못했습니다"
            ),
            axis=1,
        )
        return frame

    def _issue_frame(self, result: InspectionResult) -> pd.DataFrame:
        rows = [
            {
                "심각도": issue.severity,
                "확인 유형": issue.issue_type,
                "안내 문구": issue.message,
                "파일명": issue.file_name,
                "상세": issue.details,
            }
            for issue in result.issues
        ]
        return pd.DataFrame(rows, columns=["심각도", "확인 유형", "안내 문구", "파일명", "상세"])

    def _style_workbook(self, workbook) -> None:
        header_fill = PatternFill("solid", fgColor="1F6F5F")
        header_font = Font(color="FFFFFF", bold=True)
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            if sheet.max_row > 1 and sheet.max_column > 0:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in sheet.columns:
                max_length = 0
                column_letter = get_column_letter(column_cells[0].column)
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, min(len(value), 48))
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                sheet.column_dimensions[column_letter].width = max(10, min(max_length + 2, 50))
